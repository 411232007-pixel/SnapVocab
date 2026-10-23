from io import BytesIO

import pytest
from openpyxl import load_workbook

from snapvocab.export import ExportError, build_exports, export_filename
from snapvocab.export.xlsx import SHEET_LOG, SHEET_VOCAB, _log_titles, _vocab_titles, verify_workbook

from .conftest import make_card

ALICE, BOB = "111", "222"
BIG = 10 * 1024 * 1024


def _book(data: bytes):
    return load_workbook(BytesIO(data))


def test_two_sheets_rows_and_dedup(db):
    for _ in range(2):
        db.add_record(ALICE, make_card())
    db.add_record(ALICE, make_card(zh_name="馬克杯", en_word="mug", ja_word="マグカップ",
                                   ja_kana="マグカップ", ja_romaji="magukappu"))
    [data] = build_exports(db.all_records(ALICE), BIG, 4)
    wb = _book(data)
    assert wb.sheetnames == [SHEET_VOCAB, SHEET_LOG]
    vocab, log = wb[SHEET_VOCAB], wb[SHEET_LOG]
    assert vocab.max_row == 3  # header + 2 distinct words
    assert log.max_row == 4  # header + 3 queries
    rows = {r[1]: r for r in vocab.iter_rows(min_row=2, values_only=True)}
    assert rows["umbrella"][-1] == 2  # query count
    assert rows["mug"][-1] == 1
    assert vocab.freeze_panes == "A2"
    assert vocab.auto_filter.ref
    assert vocab["A1"].font.bold


def test_formula_like_text_is_not_a_formula(db):
    db.add_record(ALICE, make_card(en_example='=HYPERLINK("http://x","y")', note="@SUM(1)"))
    [data] = build_exports(db.all_records(ALICE), BIG, 4)
    ws = _book(data)[SHEET_VOCAB]
    for cell in ws[2]:
        assert cell.data_type != "f"
    assert ws["D2"].value.startswith("'=")


def test_cjk_text_survives(db):
    db.add_record(ALICE, make_card())
    [data] = build_exports(db.all_records(ALICE), BIG, 4)
    row = [c.value for c in _book(data)[SHEET_VOCAB][2]]
    assert "雨傘" in row and "傘" in row and "かさ" in row


def test_each_user_only_gets_their_own_rows(db):
    db.add_record(ALICE, make_card())
    db.add_record(BOB, make_card(zh_name="馬克杯", en_word="mug", en_example="Fill the mug."))
    [alice] = build_exports(db.all_records(ALICE), BIG, 4)
    [bob] = build_exports(db.all_records(BOB), BIG, 4)
    alice_text = str([[c.value for c in r] for r in _book(alice)[SHEET_LOG].iter_rows()])
    bob_text = str([[c.value for c in r] for r in _book(bob)[SHEET_LOG].iter_rows()])
    assert "umbrella" in alice_text and "mug" not in alice_text
    assert "mug" in bob_text and "umbrella" not in bob_text


def test_empty_is_rejected():
    with pytest.raises(ExportError):
        build_exports([], BIG, 4)


def test_oversized_export_is_split_into_files(db):
    for i in range(60):
        db.add_record(ALICE, make_card(en_word=f"word{i}", note="n" * 300))
    records = db.all_records(ALICE)
    [whole] = build_exports(records, BIG, 8)
    files = build_exports(records, int(len(whole) * 0.8), 8)
    assert len(files) > 1
    seen = [r[0] for f in files for r in _book(f)[SHEET_LOG].iter_rows(min_row=2, values_only=True)]
    assert sorted(seen) == sorted(r.id for r in records)


def test_impossible_size_limit_gives_clear_error(db):
    db.add_record(ALICE, make_card())
    with pytest.raises(ExportError):
        build_exports(db.all_records(ALICE), 100, 2)


def test_filename_is_safe_and_unique():
    a, b = export_filename("1234; rm -rf /"), export_filename("1234; rm -rf /")
    assert a != b
    assert a.startswith("snapvocab_1234_") and a.endswith(".xlsx")
    assert "/" not in a and " " not in a and ";" not in a


def test_headers_and_column_counts_match_spec(db):
    db.add_record(ALICE, make_card())
    [data] = build_exports(db.all_records(ALICE), BIG, 4)
    wb = _book(data)
    assert [c.value for c in wb[SHEET_VOCAB][1]] == _vocab_titles()
    assert [c.value for c in wb[SHEET_LOG][1]] == _log_titles()
    assert len(_vocab_titles()) == 14 and len(_log_titles()) == 13
    assert any("Asia/Taipei" in t for t in _vocab_titles() + _log_titles())


def test_verify_rejects_tampered_headers(db):
    db.add_record(ALICE, make_card())
    records = db.all_records(ALICE)
    [data] = build_exports(records, BIG, 4)
    wb = _book(data)
    wb[SHEET_VOCAB]["B1"] = "wrong"
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(ExportError):
        verify_workbook(buf.getvalue(), records)
