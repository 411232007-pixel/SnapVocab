"""Excel export built from one user's records (no model involved)."""

import re
import secrets
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..storage import Record
from ..timefmt import to_taipei

SHEET_VOCAB = "單字簿"
SHEET_LOG = "查詢紀錄"
TZ_LABEL = "（Asia/Taipei）"

_CARD_COLUMNS = [
    ("中文名稱", "zh_name", 14),
    ("英文單字", "en_word", 16),
    ("英文詞性", "en_pos", 10),
    ("英文例句", "en_example", 40),
    ("英文例句中譯", "en_example_zh", 32),
    ("日文單字", "ja_word", 16),
    ("日文假名", "ja_kana", 16),
    ("日文羅馬字", "ja_romaji", 18),
    ("日文例句", "ja_example", 40),
    ("日文例句中譯", "ja_example_zh", 32),
    ("語意說明", "note", 50),
]
_WRAP = {"en_example", "en_example_zh", "ja_example", "ja_example_zh", "note"}
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


class ExportError(Exception):
    """User-presentable export failure."""


def safe_filename_part(text: str, max_len: int = 40) -> str:
    """Keep only ASCII letters, digits, '-' and '_' so no path separator, '..', drive colon or
    reserved character can survive. Anything not text-safe (or empty) falls back to 'unknown'."""
    cleaned = re.sub(r"[^A-Za-z0-9_-]", "", text)[:max_len]
    return cleaned or "unknown"


def export_filename(user_id: str, part: int = 0, parts: int = 1) -> str:
    stamp = to_taipei(datetime.now().astimezone().isoformat()).strftime("%Y%m%d-%H%M%S")
    safe_id = safe_filename_part(user_id)
    suffix = f"_part{part + 1}of{parts}" if parts > 1 else ""
    return f"snapvocab_{safe_id}_{stamp}_{secrets.token_hex(3)}{suffix}.xlsx"


def _text(ws, row: int, col: int, value: str, wrap: bool = False) -> None:
    """Write a value as a literal string; neutralise anything Excel could read as a formula."""
    if value.startswith(_FORMULA_STARTS):
        value = "'" + value
    cell = ws.cell(row=row, column=col, value=value)
    cell.data_type = "s"
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _vocab_titles() -> list[str]:
    return [c[0] for c in _CARD_COLUMNS] + [f"首次查詢時間{TZ_LABEL}", f"最近查詢時間{TZ_LABEL}", "查詢次數"]


def _log_titles() -> list[str]:
    return ["紀錄 ID", f"查詢時間{TZ_LABEL}"] + [c[0] for c in _CARD_COLUMNS]


def _header(ws, titles: list[str], widths: list[int]) -> None:
    for i, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}1"


def _fmt(iso: str) -> str:
    return to_taipei(iso).strftime("%Y-%m-%d %H:%M:%S")


def _vocab_rows(records: list[Record]) -> list[dict]:
    groups: dict[tuple[str, str], list[Record]] = {}
    for r in records:
        groups.setdefault((r.en_key, r.ja_key), []).append(r)
    rows = []
    for group in groups.values():
        group.sort(key=lambda r: r.id)
        rows.append({"latest": group[-1], "first": group[0].created_at,
                     "last": group[-1].created_at, "count": len(group)})
    rows.sort(key=lambda g: g["first"])
    return rows


def _build(records: list[Record]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_VOCAB
    _header(ws, _vocab_titles(), [c[2] for c in _CARD_COLUMNS] + [22, 22, 10])
    for i, g in enumerate(_vocab_rows(records), start=2):
        card = g["latest"].card.as_dict()
        for j, (_, key, _) in enumerate(_CARD_COLUMNS, start=1):
            _text(ws, i, j, card[key], key in _WRAP)
        n = len(_CARD_COLUMNS)
        _text(ws, i, n + 1, _fmt(g["first"]))
        _text(ws, i, n + 2, _fmt(g["last"]))
        ws.cell(row=i, column=n + 3, value=g["count"])

    log = wb.create_sheet(SHEET_LOG)
    _header(log, _log_titles(), [10, 22] + [c[2] for c in _CARD_COLUMNS])
    for i, r in enumerate(sorted(records, key=lambda r: r.id), start=2):
        log.cell(row=i, column=1, value=r.id)
        _text(log, i, 2, _fmt(r.created_at))
        card = r.card.as_dict()
        for j, (_, key, _) in enumerate(_CARD_COLUMNS, start=3):
            _text(log, i, j, card[key], key in _WRAP)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def verify_workbook(data: bytes, records: list[Record]) -> None:
    """Re-open the produced file and check sheets, headers, column counts, row counts and record ids."""
    wb = load_workbook(BytesIO(data))
    if wb.sheetnames != [SHEET_VOCAB, SHEET_LOG]:
        raise ExportError("匯出檔工作表不符")
    log = wb[SHEET_LOG]
    for ws, titles in ((wb[SHEET_VOCAB], _vocab_titles()), (log, _log_titles())):
        if [c.value for c in ws[1]] != titles:
            raise ExportError("匯出檔欄位與預期不符")
        if any(len(row) != len(titles) for row in ws.iter_rows(min_row=2, values_only=True)):
            raise ExportError("匯出檔欄數與預期不符")
    ids = [row[0] for row in log.iter_rows(min_row=2, values_only=True)]
    if sorted(ids) != sorted(r.id for r in records):
        raise ExportError("匯出檔紀錄與預期不符")
    if wb[SHEET_VOCAB].max_row - 1 != len(_vocab_rows(records)):
        raise ExportError("匯出檔單字列數不符")


def build_exports(records: list[Record], size_limit: int, max_files: int) -> list[bytes]:
    """Return one or more verified .xlsx payloads, each at most size_limit bytes."""
    if not records:
        raise ExportError("尚無紀錄")
    ordered = sorted(records, key=lambda r: r.id)
    parts = 1
    while parts <= max_files:
        size = -(-len(ordered) // parts)
        chunks = [ordered[i:i + size] for i in range(0, len(ordered), size)]
        files = [_build(c) for c in chunks]
        if all(len(f) <= size_limit for f in files):
            for f, c in zip(files, chunks):
                verify_workbook(f, c)
            return files
        parts *= 2
    raise ExportError("紀錄太多，超過 Discord 附件大小上限，無法匯出")
